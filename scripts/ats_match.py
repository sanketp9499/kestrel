"""Score a tailored resume against the posting it is going to.

Every ATS on the receiving end runs some version of this comparison before a
human sees the file, and nothing in this pipeline ran it before sending.
`daily_auto_apply.relevance_score()` scores the JOB against a fixed list of
words Sanket cares about; it never reads the resume. `tailor_resume.py` injects
up to five keywords into the summary and never checks the result.

So this answers one question: of the things this posting actually asks for, how
many does the resume say? Below the threshold, phase 4 gets the specific missing
terms back and can put the true ones in rather than guessing.

Only terms that appear in the posting are ever suggested, and only Sanket's real
skills should be added from them - the missing list is a prompt to check, not a
licence to claim.
"""
import re
import unicodedata

# Phrases that appear in nearly every posting and describe no skill at all.
# Matching on them would inflate the score for free.
BOILERPLATE = {
    "equal opportunity", "equal opportunity employer", "competitive salary",
    "fast paced", "fast-paced", "about us", "benefits", "we offer", "nice to have",
    "we are looking", "looking for", "you will", "join us", "our team", "the team",
    "work closely", "close collaboration", "strong communication", "team player",
    "years experience", "years of experience", "reasonable accommodation",
    "background check", "full time", "part time", "the role", "this role",
    "the company", "our company", "day to day", "day-to-day",
}

# Words that survived the first pass against the real application folders and
# were reported as "missing skills": filler, courtesy and JD scaffolding. A
# resume is not worse for lacking them.
GENERIC = {
    "about", "make", "makes", "making", "someone", "everyone", "anyone",
    "clear", "real", "really", "environment", "benefit", "benefits", "love",
    "great", "want", "wants", "need", "needs", "join", "joining", "hiring",
    "hire", "here", "there", "thing", "things", "stuff", "lot", "lots",
    "please", "apply", "application", "applications", "candidate", "candidates",
    "position", "positions", "job", "jobs", "career", "careers", "offer",
    "offers", "offering", "life", "world", "best", "better", "right", "level",
    "part", "parts", "time", "times", "place", "today", "future", "always",
    "never", "many", "much", "come", "comes", "know", "knows", "think",
    "thinks", "take", "takes", "give", "gives", "look", "looks", "feel",
    "feels", "keep", "keeps", "value", "values", "culture", "mission",
    "vision", "passion", "passionate", "committed", "commitment", "diverse",
    "diversity", "inclusion", "inclusive", "welcome", "welcoming",
    "rather", "bring", "brings", "whether", "being", "believe", "believes",
    "learn", "learns", "learning", "first", "second", "third", "applicable",
    "ensure", "ensures", "ensuring", "provide", "provides", "providing",
    "support", "supports", "supporting", "various", "several", "multiple",
    "different", "overall", "general", "specific", "related", "relevant",
    "based", "along", "among", "toward", "towards", "around", "through",
    "throughout", "however", "therefore", "though", "although", "because",
    "employees", "employee", "staff", "member", "members", "everyone's",
    "compensation", "equity", "bonus", "insurance", "vacation", "remote",
    "hybrid", "onsite", "office", "location", "posting", "listing",
}

# The frequency pass is language-blind, so a French posting yielded "nous",
# "pour" and "avoir" as missing skills. These are the function words of the
# languages Canadian postings actually appear in.
NON_ENGLISH = {
    "nous", "vous", "pour", "avoir", "avec", "dans", "notre", "votre", "leur",
    "etre", "être", "sont", "cette", "ces", "les", "des", "une", "qui", "que",
    "plus", "sur", "par", "aux", "est", "son", "ses", "chez", "ainsi", "tout",
    "tous", "toute", "toutes", "afin", "entre", "sera", "seront", "doit",
    "devez", "recherchons", "experience", "equipe", "équipe", "poste",
    "travail", "emploi", "offrons", "avantages", "competences", "compétences",
}

STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "but", "by", "for", "from", "have",
    "has", "in", "into", "is", "it", "its", "of", "on", "or", "our", "that", "the",
    "their", "them", "they", "this", "to", "up", "us", "we", "will", "with", "you",
    "your", "who", "what", "how", "all", "any", "can", "may", "must", "such", "also",
    "across", "within", "while", "when", "where", "would", "should", "like", "well",
    "more", "most", "other", "others", "both", "each", "every", "own", "new", "help",
    "work", "working", "working with", "role", "team", "teams", "years", "year",
    "experience", "experienced", "skills", "ability", "able", "strong", "great",
    "good", "excellent", "plus", "etc", "including", "include", "includes", "using",
    "use", "used", "build", "building", "built", "create", "creating", "run", "own",
    "maintain", "looking", "require", "required", "requirements", "responsibilities",
    "opportunity", "employer", "salary", "company", "companies", "product",
    "products", "business", "customer", "customers", "user", "users", "people",
    "day", "days", "end", "closely", "deeply", "fully",
}

# Multi-word skills worth catching as a unit. A JD that says "design system"
# is not satisfied by a resume that happens to say "design" somewhere else.
PHRASES = [
    "design system", "design systems", "design token", "design tokens",
    "usability testing", "user research", "ux research", "user testing",
    "interaction design", "visual design", "product design", "service design",
    "information architecture", "user flows", "user journey", "journey mapping",
    "wireframing", "wireframes", "prototyping", "prototype", "prototypes",
    "high fidelity", "low fidelity", "responsive design", "responsive layouts",
    "accessibility", "wcag", "design thinking", "a/b testing", "ab testing",
    "motion design", "micro interactions", "style guide", "component library",
    "end to end", "end-to-end", "cross functional", "cross-functional",
    "user centred", "user centered", "user-centred", "user-centered",
    "usability", "personas", "card sorting", "heuristic evaluation",
    "design ops", "designops", "mobile design", "web app", "saas",
]

# Single-word tools and hard skills. Present or absent, no interpretation.
TOOLS = [
    "figma", "sketch", "adobe", "photoshop", "illustrator", "indesign", "xd",
    "framer", "webflow", "miro", "figjam", "invision", "principle", "protopie",
    "zeplin", "jira", "confluence", "notion", "html", "css", "javascript",
    "react", "typescript", "tailwind", "bootstrap", "wordpress", "storybook",
    "analytics", "hotjar", "amplitude", "mixpanel", "maze", "usertesting",
    "accessibility", "aria", "seo", "cms", "api", "agile", "scrum", "kanban",
]


def _normalise(text: str) -> str:
    """Lowercase, fold accents, keep the characters skills are written with.

    Folding rather than stripping matters: dropping the accent from "equipe"
    used to leave "quipe", which matched no filter and was reported as a missing
    skill on the French Crakmedia posting.
    """
    text = (text or "").lower()
    text = text.replace("’", "'").replace("–", "-").replace("—", "-")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9+/#'\- ]+", " ", text)


def _singular(term: str) -> str:
    """Crude but predictable: prototypes -> prototype, systems -> system."""
    out = []
    for word in term.split():
        if len(word) > 3 and word.endswith("s") and not word.endswith("ss"):
            word = word[:-1]
        out.append(word)
    return " ".join(out)


def extract_keywords(jd_text: str, limit: int = 25, company: str = "") -> list:
    """The terms this posting actually asks for, most important first.

    Phrases and tools are taken as-is when present. A few remaining slots go to
    the most frequent meaningful words, which is what picks up domain language
    ("payments", "onboarding") that no fixed vocabulary can predict.

    The frequency pass is deliberately the smaller half. Run unfiltered over the
    real application folders it produced "crakmedia", "altaml", "about" and
    "someone" as missing skills, which is noise the gate would then send phase 4
    chasing.
    """
    text = _normalise(jd_text)
    if not text.strip():
        return []

    company_words = {w for w in _normalise(company).split() if len(w) > 2}

    found, seen = [], set()

    def add(term):
        key = _singular(term)
        if key in seen:
            return
        seen.add(key)
        found.append(term)

    for phrase in PHRASES:
        if phrase in text and phrase not in BOILERPLATE:
            add(phrase)
    for tool in TOOLS:
        if re.search(r"\b%s\b" % re.escape(tool), text):
            add(tool)

    curated = len(found)
    freq_cap = max(3, limit // 3)  # curated vocabulary carries the signal

    counts = {}
    for word in text.split():
        word = word.strip("-'/")
        if len(word) < 5 or not word[0].isalpha():
            continue
        if "'" in word:  # we're, you'll, l'experience
            continue
        if (word in STOPWORDS or word in GENERIC or word in NON_ENGLISH
                or word in BOILERPLATE or word in company_words):
            continue
        counts[word] = counts.get(word, 0) + 1
    for word, n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
        if n < 3 or len(found) - curated >= freq_cap or len(found) >= limit:
            continue
        if any(word in f for f in found):
            continue
        add(word)

    return found[:limit]


def coverage(jd_text: str, resume_text: str, company: str = "") -> dict:
    """How much of the posting's language the resume actually contains.

    Returns score (0.0-1.0), the keywords considered, and which were matched or
    missed. An empty JD scores 1.0: a failed scrape must not block a real
    application.
    """
    keywords = extract_keywords(jd_text, company=company)
    if not keywords:
        return {"score": 1.0, "keywords": [], "matched": [], "missing": []}

    haystack = _singular(_normalise(resume_text))
    matched, missing = [], []
    for kw in keywords:
        (matched if _singular(kw) in haystack else missing).append(kw)

    return {"score": round(len(matched) / len(keywords), 3),
            "keywords": keywords, "matched": matched, "missing": missing}


def gate(jd_text: str, resume_text: str, threshold: float = 0.70,
         company: str = "") -> tuple:
    """(passed, result). Below the threshold, result["missing"] says what to fix."""
    result = coverage(jd_text, resume_text, company=company)
    return result["score"] >= threshold, result


def resume_text(path: str) -> str:
    """Plain text of a .docx or .txt resume, for scoring the file that will
    actually be sent rather than the text it was built from."""
    if path.lower().endswith(".docx"):
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    import io
    return io.open(path, encoding="utf-8", errors="replace").read()


if __name__ == "__main__":
    import argparse, io, json
    ap = argparse.ArgumentParser(description="Score a resume against a job description")
    ap.add_argument("--jd-file", required=True, help="Job_Description.md for the posting")
    ap.add_argument("--resume", required=True, help=".docx or .txt resume to score")
    ap.add_argument("--threshold", type=float, default=0.70)
    ap.add_argument("--company", default="", help="excluded from keywords")
    args = ap.parse_args()

    jd = io.open(args.jd_file, encoding="utf-8", errors="replace").read()
    ok, result = gate(jd, resume_text(args.resume), args.threshold, args.company)
    print(json.dumps({"passed": ok, **result}, indent=2, ensure_ascii=False))
    raise SystemExit(0 if ok else 1)
